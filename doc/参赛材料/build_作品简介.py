from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from reportlab.lib.colors import HexColor
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parent
TITLE = "智航灵枢：面向复杂任务的多模态具身智能无人机系统"
BODY = (
    "传统无人机地面站操作门槛高，直接使用大模型则存在幻觉、越权和结果难验证等风险。"
    "智航灵枢构建多模态具身智能无人机系统：LLM融合飞控遥测、RGB-D、点云、YOLO与VLM证据，"
    "将开放指令转化为受约束、可确认的任务流程；确定性网关校验参数并触发高风险确认，"
    "ROS 2/PX4执行飞行，遥测验证物理终态。系统已实现跨端交互、航点飞行、避障重规划、拍照分析和任务报告。"
    "80条任务四组对比中，混合Agent严格通过率57.5%，约为纯规则的2.7倍；"
    "在AirSim/PX4仿真中，固定任务10次闭环成功率100%。"
    "项目面向巡检、搜救、园区安防与科研教学。"
)
COUNT = len(BODY)
assert COUNT <= 300, f"正文超过 300 字：{COUNT}"
PDF_FONT = "STSong-Light"


def register_pdf_font():
    global PDF_FONT
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path("/mnt/c/Windows/Fonts/msyh.ttc"),
        Path("/mnt/c/Windows/Fonts/NotoSansSC-VF.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont("AeroMindCJK", str(candidate), subfontIndex=0))
            PDF_FONT = "AeroMindCJK"
            return
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


def set_run_font(run, name: str, size: float, bold: bool = False, color=None):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    if color:
        run.font.color.rgb = RGBColor(*color)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)


def set_cell_shading(paragraph, fill: str):
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    p_pr.append(shd)


def build_docx():
    doc = Document()
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.4)
    section.left_margin = Cm(2.6)
    section.right_margin = Cm(2.6)

    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(18)
    run = title.add_run(TITLE)
    set_run_font(run, "微软雅黑", 20, True, (18, 48, 78))

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(20)
    run = label.add_run("参赛作品简介")
    set_run_font(run, "微软雅黑", 11, False, (53, 132, 126))

    body = doc.add_paragraph()
    body.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    body.paragraph_format.first_line_indent = Cm(0.74)
    body.paragraph_format.line_spacing = 1.65
    body.paragraph_format.space_after = Pt(18)
    run = body.add_run(BODY)
    set_run_font(run, "宋体", 12)

    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    note.paragraph_format.space_before = Pt(8)
    set_cell_shading(note, "EAF3F2")
    run = note.add_run(f"正文 {COUNT} 字｜不含标题")
    set_run_font(run, "微软雅黑", 9, False, (53, 90, 92))

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer.add_run("第八届中国研究生人工智能创新大赛")
    set_run_font(run, "微软雅黑", 8, False, (110, 118, 126))

    doc.save(ROOT / "智航灵枢_参赛作品简介.docx")


def build_pdf():
    register_pdf_font()
    out = ROOT / "智航灵枢_参赛作品简介.pdf"
    doc = SimpleDocTemplate(
        str(out),
        pagesize=A4,
        rightMargin=2.6 * cm,
        leftMargin=2.6 * cm,
        topMargin=2.8 * cm,
        bottomMargin=2.5 * cm,
        title=TITLE,
        author="智航灵枢项目团队",
    )
    title_style = ParagraphStyle(
        "TitleCN",
        fontName=PDF_FONT,
        fontSize=20,
        leading=29,
        alignment=TA_CENTER,
        textColor=HexColor("#12304E"),
        spaceAfter=13,
    )
    label_style = ParagraphStyle(
        "LabelCN",
        fontName=PDF_FONT,
        fontSize=11,
        leading=16,
        alignment=TA_CENTER,
        textColor=HexColor("#35847E"),
        spaceAfter=24,
    )
    body_style = ParagraphStyle(
        "BodyCN",
        fontName=PDF_FONT,
        fontSize=12,
        leading=24,
        alignment=TA_JUSTIFY,
        firstLineIndent=24,
        textColor=HexColor("#202830"),
    )
    note_style = ParagraphStyle(
        "NoteCN",
        fontName=PDF_FONT,
        fontSize=9,
        leading=14,
        alignment=TA_CENTER,
        textColor=HexColor("#35605C"),
        backColor=HexColor("#EAF3F2"),
        borderPadding=(7, 8, 7, 8),
        spaceBefore=20,
    )
    story = [
        Paragraph(TITLE, title_style),
        Paragraph("参赛作品简介", label_style),
        Spacer(1, 0.25 * cm),
        Paragraph(BODY, body_style),
        Paragraph(f"正文 {COUNT} 字　｜　不含标题", note_style),
    ]
    doc.build(story)


if __name__ == "__main__":
    build_docx()
    build_pdf()
    print(f"已生成 DOCX/PDF，正文 {COUNT} 字")
