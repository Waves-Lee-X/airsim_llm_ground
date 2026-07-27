from __future__ import annotations

from dataclasses import dataclass
from datetime import datetime, timezone
from html import escape
from pathlib import Path
import re

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from PIL import Image as PILImage
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER, TA_JUSTIFY, TA_LEFT
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import (
    Image,
    LongTable,
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)


ROOT = Path(__file__).resolve().parent
SOURCE = ROOT / "翼策_项目文档.md"
DOCX_OUT = ROOT / "翼策_项目文档.docx"
PDF_OUT = ROOT / "翼策_项目文档.pdf"

PROJECT = "翼策——面向复杂低空任务的多模态具身智能无人机系统"
VERSION = "V2.0 正式版"
DATE = "2026.07.24"
TEAM = "翼策项目团队"
GROUP = "开放赛题（赛题四：机器人与具身智能）"
STAGE = "AirSim/PX4 仿真研究原型"
DOCUMENT_KIND = "初赛项目文档正式稿"

NAVY = "173A5E"
TEAL = "1E7A78"
PALE = "EAF2F4"
LIGHT = "F5F7F9"
TEXT = "26333D"
MUTED = "647481"
PDF_FONT = "STSong-Light"


def register_pdf_font():
    global PDF_FONT
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path("/mnt/c/Windows/Fonts/msyh.ttc"),
        Path("/mnt/c/Windows/Fonts/NotoSansSC-VF.ttf"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    ]
    for candidate in candidates:
        if candidate.exists():
            pdfmetrics.registerFont(TTFont("AeroMindCJK", str(candidate), subfontIndex=0))
            PDF_FONT = "AeroMindCJK"
            return
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))


@dataclass
class Block:
    kind: str
    value: object
    level: int = 0


def normalize_inline(text: str) -> str:
    text = re.sub(r"!\[([^\]]*)\]\([^)]+\)", r"\1", text)
    text = re.sub(r"\[([^\]]+)\]\(([^)]+)\)", r"\1（\2）", text)
    text = text.replace("**", "").replace("__", "").replace("`", "")
    return text.strip()


def parse_blocks(markdown: str) -> list[Block]:
    lines = markdown.splitlines()
    blocks: list[Block] = []
    paragraph: list[str] = []
    code: list[str] = []
    in_code = False
    index = 0

    def flush_paragraph():
        if paragraph:
            blocks.append(Block("paragraph", normalize_inline("".join(paragraph))))
            paragraph.clear()

    while index < len(lines):
        raw = lines[index]
        stripped = raw.strip()
        if stripped.startswith("```"):
            if in_code:
                blocks.append(Block("code", "\n".join(code)))
                code.clear()
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            index += 1
            continue
        if in_code:
            code.append(raw)
            index += 1
            continue
        if not stripped:
            flush_paragraph()
            index += 1
            continue
        image_match = re.fullmatch(r"!\[([^\]]*)\]\(([^)]+)\)", stripped)
        if image_match:
            flush_paragraph()
            blocks.append(
                Block(
                    "image",
                    {
                        "caption": image_match.group(1),
                        "path": (ROOT / image_match.group(2)).resolve(),
                    },
                )
            )
            index += 1
            continue
        heading_match = re.match(r"^(#{1,6})\s+(.+)$", stripped)
        if heading_match:
            flush_paragraph()
            blocks.append(
                Block(
                    "heading",
                    normalize_inline(heading_match.group(2)),
                    len(heading_match.group(1)),
                )
            )
            index += 1
            continue
        if stripped.startswith("|") and index + 1 < len(lines):
            separator = lines[index + 1].strip()
            if separator.startswith("|") and "---" in separator:
                flush_paragraph()
                rows = []
                while index < len(lines) and lines[index].strip().startswith("|"):
                    cells = [
                        normalize_inline(cell)
                        for cell in lines[index].strip().strip("|").split("|")
                    ]
                    if index == 0 or not all(
                        re.fullmatch(r":?-{3,}:?", cell.replace(" ", ""))
                        for cell in cells
                    ):
                        rows.append(cells)
                    index += 1
                    if len(rows) == 1 and index < len(lines):
                        index += 1
                blocks.append(Block("table", rows))
                continue
        if stripped.startswith(">"):
            flush_paragraph()
            blocks.append(Block("quote", normalize_inline(stripped.lstrip("> "))))
            index += 1
            continue
        list_match = re.match(r"^([-*]|\d+\.)\s+(.+)$", stripped)
        if list_match:
            flush_paragraph()
            kind = "number" if list_match.group(1)[0].isdigit() else "bullet"
            blocks.append(Block(kind, normalize_inline(list_match.group(2))))
            index += 1
            continue
        paragraph.append(stripped)
        index += 1
    flush_paragraph()
    if code:
        blocks.append(Block("code", "\n".join(code)))
    return blocks


def content_blocks() -> list[Block]:
    blocks = parse_blocks(SOURCE.read_text(encoding="utf-8"))
    start = next(
        index
        for index, block in enumerate(blocks)
        if block.kind == "heading" and block.value == "记录更改历史"
    )
    return blocks[start:]


def set_docx_font(run, name="宋体", size=10.5, bold=False, color=TEXT):
    run.font.name = name
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), name)


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=80, start=90, bottom=80, end=90):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def add_docx_page_number(paragraph):
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    run._r.extend([fld_char1, instr_text, fld_char2])
    set_docx_font(run, "微软雅黑", 8, False, MUTED)


def add_docx_toc(paragraph):
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = 'TOC \\o "1-3" \\h \\z \\u'
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "打开 Word 后右键此处并选择“更新域”，即可生成目录。"
    separate.append(text)
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend([begin, instruction, separate, end])
    set_docx_font(run, "宋体", 10.5, False, MUTED)


def configure_docx_styles(doc: Document):
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "宋体"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.paragraph_format.line_spacing = 1.45
    normal.paragraph_format.space_after = Pt(5)
    for name, font, size, color in (
        ("Title", "微软雅黑", 24, NAVY),
        ("Heading 1", "微软雅黑", 16, NAVY),
        ("Heading 2", "微软雅黑", 13, TEAL),
        ("Heading 3", "微软雅黑", 11, NAVY),
        ("Heading 4", "微软雅黑", 10.5, NAVY),
    ):
        style = styles[name]
        style.font.name = font
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:eastAsia"), font)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.space_before = Pt(12 if name != "Heading 1" else 18)
        style.paragraph_format.space_after = Pt(6)


def add_docx_cover(doc: Document):
    section = doc.sections[0]
    section.page_height = Cm(29.7)
    section.page_width = Cm(21)
    section.top_margin = Cm(2.4)
    section.bottom_margin = Cm(2.2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(2.5)

    spacer = doc.add_paragraph()
    spacer.paragraph_format.space_after = Pt(42)
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title.paragraph_format.space_after = Pt(28)
    run = title.add_run("第八届中国研究生人工智能创新大赛")
    set_docx_font(run, "微软雅黑", 17, True, NAVY)

    project = doc.add_paragraph()
    project.alignment = WD_ALIGN_PARAGRAPH.CENTER
    project.paragraph_format.space_after = Pt(18)
    run = project.add_run(PROJECT)
    set_docx_font(run, "微软雅黑", 23, True, NAVY)

    label = doc.add_paragraph()
    label.alignment = WD_ALIGN_PARAGRAPH.CENTER
    label.paragraph_format.space_after = Pt(58)
    run = label.add_run("项目文档")
    set_docx_font(run, "微软雅黑", 15, False, TEAL)

    table = doc.add_table(rows=6, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    values = [
        ("版本", VERSION),
        ("日期", DATE),
        ("团队名称", TEAM),
        ("参赛组别", GROUP),
        ("当前阶段", STAGE),
        ("文档性质", DOCUMENT_KIND),
    ]
    for row, (key, value) in zip(table.rows, values):
        row.cells[0].width = Cm(4.0)
        row.cells[1].width = Cm(10.0)
        shade_cell(row.cells[0], PALE)
        for cell, text in zip(row.cells, (key, value)):
            set_cell_margins(cell, 100, 120, 100, 120)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_docx_font(p.add_run(text), "微软雅黑", 10.5, cell == row.cells[0], TEXT)
    note = doc.add_paragraph()
    note.alignment = WD_ALIGN_PARAGRAPH.CENTER
    note.paragraph_format.space_before = Pt(36)
    set_docx_font(
        note.add_run("开放赛题四｜AirSim/PX4 仿真验证｜V2.0 正式版"),
        "微软雅黑",
        9,
        False,
        MUTED,
    )
    doc.add_page_break()


def add_docx_blocks(doc: Document, blocks: list[Block]):
    figure_index = 0
    for block in blocks:
        if block.kind == "heading":
            level = max(1, min(block.level, 4))
            doc.add_heading(str(block.value), level=level)
        elif block.kind == "paragraph":
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            p.paragraph_format.first_line_indent = Cm(0.74)
            set_docx_font(p.add_run(str(block.value)))
        elif block.kind == "quote":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.7)
            p.paragraph_format.right_indent = Cm(0.7)
            p.paragraph_format.space_before = Pt(5)
            p.paragraph_format.space_after = Pt(8)
            set_docx_font(p.add_run(str(block.value)), "微软雅黑", 10, False, TEAL)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), PALE)
            p_pr.append(shd)
        elif block.kind in {"bullet", "number"}:
            style = "List Bullet" if block.kind == "bullet" else "List Number"
            p = doc.add_paragraph(style=style)
            p.paragraph_format.space_after = Pt(2)
            set_docx_font(p.add_run(str(block.value)))
        elif block.kind == "code":
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.right_indent = Cm(0.6)
            p.paragraph_format.space_before = Pt(4)
            p.paragraph_format.space_after = Pt(8)
            p_pr = p._p.get_or_add_pPr()
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), "EEF2F5")
            p_pr.append(shd)
            set_docx_font(p.add_run(str(block.value)), "等线", 9, False, TEXT)
        elif block.kind == "table":
            rows = block.value
            if not rows:
                continue
            width = max(len(row) for row in rows)
            table = doc.add_table(rows=len(rows), cols=width)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER
            table.style = "Table Grid"
            for row_index, row in enumerate(rows):
                for col_index in range(width):
                    cell = table.cell(row_index, col_index)
                    value = row[col_index] if col_index < len(row) else ""
                    set_cell_margins(cell)
                    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
                    if row_index == 0:
                        shade_cell(cell, NAVY)
                    elif row_index % 2 == 0:
                        shade_cell(cell, LIGHT)
                    p = cell.paragraphs[0]
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER if row_index == 0 else WD_ALIGN_PARAGRAPH.LEFT
                    set_docx_font(
                        p.add_run(value),
                        "微软雅黑" if row_index == 0 else "宋体",
                        8 if width >= 7 else 9,
                        row_index == 0,
                        "FFFFFF" if row_index == 0 else TEXT,
                    )
            doc.add_paragraph().paragraph_format.space_after = Pt(2)
        elif block.kind == "image":
            path = block.value["path"]
            if not path.exists():
                p = doc.add_paragraph()
                set_docx_font(p.add_run(f"[图片缺失：{path}]"), color="A33A3A")
                continue
            figure_index += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run()
            run.add_picture(str(path), width=Cm(15.6))
            caption = doc.add_paragraph()
            caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
            set_docx_font(
                caption.add_run(f"图 {figure_index}  {block.value['caption']}"),
                "微软雅黑",
                9,
                False,
                MUTED,
            )


def build_docx(blocks: list[Block]):
    doc = Document()
    props = doc.core_properties
    props.title = PROJECT
    props.subject = "第八届中国研究生人工智能创新大赛项目文档"
    props.author = TEAM
    props.last_modified_by = TEAM
    props.category = GROUP
    props.keywords = "多模态具身智能, 无人机, ROS 2, PX4, Agent"
    props.comments = "V2.0 正式版"
    props.created = datetime(2026, 7, 24, tzinfo=timezone.utc)
    props.modified = datetime(2026, 7, 24, tzinfo=timezone.utc)
    configure_docx_styles(doc)
    add_docx_cover(doc)

    toc = doc.add_heading("目录", level=1)
    toc.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_docx_toc(doc.add_paragraph())
    doc.add_page_break()

    section = doc.sections[-1]
    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    set_docx_font(header.add_run("翼策｜项目文档"), "微软雅黑", 8, False, MUTED)
    add_docx_page_number(section.footer.paragraphs[0])
    add_docx_blocks(doc, blocks)

    settings = doc.settings._element
    update_fields = OxmlElement("w:updateFields")
    update_fields.set(qn("w:val"), "true")
    settings.append(update_fields)
    doc.save(DOCX_OUT)


def pdf_text(text: str) -> str:
    text = normalize_inline(text)
    return escape(text).replace("\n", "<br/>")


def pdf_styles():
    base = getSampleStyleSheet()
    return {
        "title": ParagraphStyle(
            "CNTitle",
            parent=base["Title"],
            fontName=PDF_FONT,
            fontSize=22,
            leading=32,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceAfter=18,
            wordWrap="CJK",
        ),
        "subtitle": ParagraphStyle(
            "CNSubtitle",
            fontName=PDF_FONT,
            fontSize=15,
            leading=22,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{TEAL}"),
            spaceAfter=18,
            wordWrap="CJK",
        ),
        "h1": ParagraphStyle(
            "CNH1",
            fontName=PDF_FONT,
            fontSize=15,
            leading=22,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=14,
            spaceAfter=8,
            keepWithNext=True,
            wordWrap="CJK",
        ),
        "h2": ParagraphStyle(
            "CNH2",
            fontName=PDF_FONT,
            fontSize=12.5,
            leading=19,
            textColor=colors.HexColor(f"#{TEAL}"),
            spaceBefore=11,
            spaceAfter=6,
            keepWithNext=True,
            wordWrap="CJK",
        ),
        "h3": ParagraphStyle(
            "CNH3",
            fontName=PDF_FONT,
            fontSize=11,
            leading=17,
            textColor=colors.HexColor(f"#{NAVY}"),
            spaceBefore=8,
            spaceAfter=4,
            keepWithNext=True,
            wordWrap="CJK",
        ),
        "body": ParagraphStyle(
            "CNBody",
            fontName=PDF_FONT,
            fontSize=9.5,
            leading=16,
            alignment=TA_JUSTIFY,
            firstLineIndent=19,
            textColor=colors.HexColor(f"#{TEXT}"),
            spaceAfter=5,
            wordWrap="CJK",
        ),
        "list": ParagraphStyle(
            "CNList",
            fontName=PDF_FONT,
            fontSize=9.5,
            leading=15,
            leftIndent=18,
            firstLineIndent=-10,
            textColor=colors.HexColor(f"#{TEXT}"),
            spaceAfter=2,
            wordWrap="CJK",
        ),
        "quote": ParagraphStyle(
            "CNQuote",
            fontName=PDF_FONT,
            fontSize=9.5,
            leading=16,
            leftIndent=14,
            rightIndent=14,
            borderColor=colors.HexColor(f"#{TEAL}"),
            borderWidth=1,
            borderPadding=7,
            backColor=colors.HexColor(f"#{PALE}"),
            textColor=colors.HexColor(f"#{TEAL}"),
            spaceBefore=5,
            spaceAfter=8,
            wordWrap="CJK",
        ),
        "code": ParagraphStyle(
            "CNCode",
            fontName=PDF_FONT,
            fontSize=8.5,
            leading=13,
            leftIndent=10,
            rightIndent=10,
            borderPadding=7,
            backColor=colors.HexColor("#EEF2F5"),
            textColor=colors.HexColor(f"#{TEXT}"),
            spaceBefore=4,
            spaceAfter=8,
            wordWrap="CJK",
        ),
        "caption": ParagraphStyle(
            "CNCaption",
            fontName=PDF_FONT,
            fontSize=8.5,
            leading=13,
            alignment=TA_CENTER,
            textColor=colors.HexColor(f"#{MUTED}"),
            spaceAfter=8,
            wordWrap="CJK",
        ),
        "table": ParagraphStyle(
            "CNTable",
            fontName=PDF_FONT,
            fontSize=7.2,
            leading=10,
            textColor=colors.HexColor(f"#{TEXT}"),
            wordWrap="CJK",
        ),
        "table_head": ParagraphStyle(
            "CNTableHead",
            fontName=PDF_FONT,
            fontSize=7.2,
            leading=10,
            alignment=TA_CENTER,
            textColor=colors.white,
            wordWrap="CJK",
        ),
    }


def scaled_image(path: Path, max_width: float, max_height: float):
    with PILImage.open(path) as image:
        width, height = image.size
    scale = min(max_width / width, max_height / height)
    return Image(str(path), width=width * scale, height=height * scale)


def pdf_table(rows, styles, available_width):
    width = max(len(row) for row in rows)
    data = []
    for row_index, row in enumerate(rows):
        data.append(
            [
                Paragraph(
                    pdf_text(row[col_index] if col_index < len(row) else ""),
                    styles["table_head"] if row_index == 0 else styles["table"],
                )
                for col_index in range(width)
            ]
        )
    col_widths = [available_width / width] * width
    table = LongTable(data, colWidths=col_widths, repeatRows=1, hAlign="CENTER")
    commands = [
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(f"#{NAVY}")),
        ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#B8C4CC")),
        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
        ("LEFTPADDING", (0, 0), (-1, -1), 4),
        ("RIGHTPADDING", (0, 0), (-1, -1), 4),
        ("TOPPADDING", (0, 0), (-1, -1), 4),
        ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
    ]
    for row_index in range(1, len(data)):
        if row_index % 2 == 0:
            commands.append(
                ("BACKGROUND", (0, row_index), (-1, row_index), colors.HexColor(f"#{LIGHT}"))
            )
    table.setStyle(TableStyle(commands))
    return table


def build_pdf(blocks: list[Block]):
    register_pdf_font()
    styles = pdf_styles()
    doc = SimpleDocTemplate(
        str(PDF_OUT),
        pagesize=A4,
        rightMargin=1.8 * cm,
        leftMargin=1.8 * cm,
        topMargin=1.8 * cm,
        bottomMargin=1.7 * cm,
        title=PROJECT,
        author=TEAM,
        subject="第八届中国研究生人工智能创新大赛项目文档",
    )
    available_width = A4[0] - 3.6 * cm
    story = [
        Spacer(1, 1.6 * cm),
        Paragraph("第八届中国研究生人工智能创新大赛", styles["subtitle"]),
        Spacer(1, 0.7 * cm),
        Paragraph(PROJECT, styles["title"]),
        Paragraph("项目文档", styles["subtitle"]),
        Spacer(1, 1.0 * cm),
    ]
    metadata = [
        ["版本", VERSION],
        ["日期", DATE],
        ["团队名称", TEAM],
        ["参赛组别", GROUP],
        ["当前阶段", STAGE],
        ["文档性质", DOCUMENT_KIND],
    ]
    cover_table = Table(
        [
            [
                Paragraph(pdf_text(row[0]), styles["table_head"]),
                Paragraph(pdf_text(row[1]), styles["table"]),
            ]
            for row in metadata
        ],
        colWidths=[4.0 * cm, 10.0 * cm],
        hAlign="CENTER",
    )
    cover_table.setStyle(
        TableStyle(
            [
                ("BACKGROUND", (0, 0), (0, -1), colors.HexColor(f"#{NAVY}")),
                ("BACKGROUND", (1, 0), (1, -1), colors.HexColor(f"#{LIGHT}")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8C4CC")),
                ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ("TOPPADDING", (0, 0), (-1, -1), 8),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 8),
            ]
        )
    )
    story.extend(
        [
            cover_table,
            Spacer(1, 1.0 * cm),
            Paragraph(
                "开放赛题四｜AirSim/PX4 仿真验证｜V2.0 正式版",
                styles["caption"],
            ),
            PageBreak(),
            Paragraph("目录", styles["title"]),
            Spacer(1, 0.5 * cm),
        ]
    )
    for line in (
        "1 项目概况",
        "　1.1 背景和基础",
        "　1.2 场景和价值",
        "　1.3 所需支持",
        "2 项目规划",
        "　2.1 整体目标",
        "　2.2 技术创新点",
        "3 实施方案",
        "　3.1 技术可行性分析",
        "　3.2 技术细节",
        "　3.3 计划和分工",
        "4 参考资料",
    ):
        story.append(Paragraph(pdf_text(line), styles["body"]))
    story.append(PageBreak())

    figure_index = 0
    number_index = 0
    for block in blocks:
        if block.kind == "heading":
            style = styles["h1"] if block.level <= 1 else styles["h2"] if block.level == 2 else styles["h3"]
            story.append(Paragraph(pdf_text(str(block.value)), style))
        elif block.kind == "paragraph":
            story.append(Paragraph(pdf_text(str(block.value)), styles["body"]))
        elif block.kind == "quote":
            story.append(Paragraph(pdf_text(str(block.value)), styles["quote"]))
        elif block.kind in {"bullet", "number"}:
            if block.kind == "number":
                number_index += 1
                prefix = f"{number_index}."
            else:
                number_index = 0
                prefix = "•"
            story.append(Paragraph(f"{prefix}　{pdf_text(str(block.value))}", styles["list"]))
        elif block.kind == "code":
            number_index = 0
            story.append(Paragraph(pdf_text(str(block.value)), styles["code"]))
        elif block.kind == "table":
            number_index = 0
            story.extend(
                [
                    pdf_table(block.value, styles, available_width),
                    Spacer(1, 0.18 * cm),
                ]
            )
        elif block.kind == "image":
            number_index = 0
            path = block.value["path"]
            if path.exists():
                figure_index += 1
                story.append(scaled_image(path, available_width, 12.5 * cm))
                story.append(
                    Paragraph(
                        f"图 {figure_index}　{pdf_text(block.value['caption'])}",
                        styles["caption"],
                    )
                )

    def page(canvas, _doc):
        canvas.saveState()
        canvas.setFont(PDF_FONT, 7.5)
        canvas.setFillColor(colors.HexColor(f"#{MUTED}"))
        canvas.drawString(1.8 * cm, 0.85 * cm, "翼策｜项目文档")
        canvas.drawRightString(A4[0] - 1.8 * cm, 0.85 * cm, str(_doc.page))
        canvas.restoreState()

    doc.build(story, onFirstPage=page, onLaterPages=page)


def main():
    blocks = content_blocks()
    build_docx(blocks)
    build_pdf(blocks)
    print(f"已生成：{DOCX_OUT}")
    print(f"已生成：{PDF_OUT}")


if __name__ == "__main__":
    main()
